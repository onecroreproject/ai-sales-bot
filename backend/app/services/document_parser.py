import io
import re
import csv
import logging
from typing import BinaryIO
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.product_knowledge import ProductKnowledge
from app.ai.embeddings import create_embedding
from app.services.crawler_service import chunk_text

logger = logging.getLogger(__name__)


def extract_text_from_pdf(file_bytes: bytes) -> str:
    """Extract clean text from PDF file bytes."""
    import pypdf
    reader = pypdf.PdfReader(io.BytesIO(file_bytes))
    pages_text = []
    for idx, page in enumerate(reader.pages):
        page_content = page.extract_text() or ""
        if page_content.strip():
            pages_text.append(f"[Page {idx + 1}]\n{page_content.strip()}")
    return "\n\n".join(pages_text)


def extract_text_from_docx(file_bytes: bytes) -> str:
    """Extract paragraphs and tables from Word (.docx) file bytes."""
    import docx
    doc = docx.Document(io.BytesIO(file_bytes))
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    
    # Also extract tables
    for table in doc.tables:
        for row in table.rows:
            row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_cells:
                paragraphs.append(" | ".join(row_cells))

    return "\n\n".join(paragraphs)


def extract_text_from_excel(file_bytes: bytes) -> str:
    """Extract sheets and rows from Excel (.xlsx) file bytes."""
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    rows_text = []

    for sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]
        headers = []
        for row_idx, row in enumerate(sheet.iter_rows(values_only=True)):
            if not any(row):
                continue
            cleaned = [str(val).strip() for val in row if val is not None]
            if not cleaned:
                continue
            if row_idx == 0:
                headers = cleaned
                rows_text.append(f"Sheet [{sheet_name}] Headers: " + ", ".join(headers))
            else:
                if headers and len(headers) == len(cleaned):
                    row_pairs = [f"{headers[i]}: {cleaned[i]}" for i in range(len(cleaned))]
                    rows_text.append(f"Sheet [{sheet_name}] Row {row_idx}: " + "; ".join(row_pairs))
                else:
                    rows_text.append(f"Sheet [{sheet_name}] Row {row_idx}: " + " | ".join(cleaned))

    return "\n".join(rows_text)


def extract_text_from_csv(file_bytes: bytes) -> str:
    """Extract structured rows from CSV bytes."""
    text_content = file_bytes.decode("utf-8", errors="ignore")
    reader = csv.reader(io.StringIO(text_content))
    rows_text = []
    headers = []

    for idx, row in enumerate(reader):
        cleaned = [c.strip() for c in row if c.strip()]
        if not cleaned:
            continue
        if idx == 0:
            headers = cleaned
            rows_text.append("CSV Columns: " + ", ".join(headers))
        else:
            if headers and len(headers) == len(cleaned):
                row_pairs = [f"{headers[i]}: {cleaned[i]}" for i in range(len(cleaned))]
                rows_text.append(f"Record {idx}: " + "; ".join(row_pairs))
            else:
                rows_text.append(f"Record {idx}: " + " | ".join(cleaned))

    return "\n".join(rows_text)


async def parse_and_index_document(
    db: AsyncSession,
    company_id: int,
    filename: str,
    file_bytes: bytes,
    product_id: int | None = None,
    api_key: str | None = None,
) -> dict:
    """
    Parses PDF, Word, Excel, CSV, or Text file, splits into semantic chunks,
    generates 1536-dimension embeddings, and stores into ProductKnowledge.
    """
    lower_name = filename.lower()
    
    if lower_name.endswith(".pdf"):
        full_text = extract_text_from_pdf(file_bytes)
    elif lower_name.endswith(".docx"):
        full_text = extract_text_from_docx(file_bytes)
    elif lower_name.endswith((".xlsx", ".xls")):
        full_text = extract_text_from_excel(file_bytes)
    elif lower_name.endswith(".csv"):
        full_text = extract_text_from_csv(file_bytes)
    elif lower_name.endswith((".txt", ".md", ".json")):
        full_text = file_bytes.decode("utf-8", errors="ignore")
    else:
        raise ValueError(f"Unsupported file format '{filename}'. Supported formats: PDF, DOCX, XLSX, CSV, TXT.")

    full_text = re.sub(r'\s+', ' ', full_text).strip()

    if len(full_text) < 20:
        raise ValueError(f"The document '{filename}' did not contain enough readable text content.")

    chunks = chunk_text(full_text, max_chunk_size=1000, overlap=120)
    created_count = 0

    for chunk in chunks:
        formatted_content = f"[File: {filename}]\n{chunk}"
        try:
            embedding = await create_embedding(formatted_content, api_key=api_key)
        except Exception as embed_err:
            logger.warning(f"Could not generate embedding for document chunk: {embed_err}. Using placeholder embedding.")
            embedding = [0.0] * 1536

        knowledge_entry = ProductKnowledge(
            company_id=company_id,
            product_id=product_id,
            content=formatted_content,
            embedding=embedding,
        )
        db.add(knowledge_entry)
        created_count += 1

    await db.commit()

    return {
        "filename": filename,
        "chunks_indexed": created_count,
        "total_characters": len(full_text),
    }
